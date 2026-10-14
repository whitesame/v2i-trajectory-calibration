"""Build the course presentation from the supplied group-meeting deck.

Requires: python-pptx. Run from this file's directory or any working directory.
The two retained group-analysis slides stay editable, including their charts.
"""

from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
SOURCE = ROOT / "基于目标轨迹的目标级车路协同定位.pptx"
OUTPUT = HERE / "现代测量数据处理_目标级车路配准.pptx"

NAVY = RGBColor(0, 30, 100)
BLUE = RGBColor(32, 61, 136)
PALE = RGBColor(235, 241, 250)
WHITE = RGBColor(255, 255, 255)
RED = RGBColor(192, 0, 0)
GRAY = RGBColor(85, 85, 85)
LIGHT = RGBColor(246, 248, 252)


def rect(slide, x, y, w, h, fill, radius=False, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background() if line is None else None
    if line is not None:
        shape.line.color.rgb = line
    return shape


def text(slide, content, x, y, w, h, size=18, color=NAVY, bold=False,
         align=PP_ALIGN.LEFT, margin=0.04):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(content.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(3)
        run = p.add_run()
        run.text = line
        run.font.name = "微软雅黑"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def base(deck, sec, title, source="V2I-Calib, IROS 2024；上游代码仓库"):
    s = deck.slides.add_slide(deck.slide_layouts[11])
    rect(s, 0, 0, 13.333, 7.5, WHITE)
    rect(s, 0, 0, 0.18, 7.5, NAVY)
    rect(s, 0.18, 0, 13.15, 0.10, NAVY)
    text(s, sec, 0.42, 0.32, 0.72, 0.55, 28, NAVY, True)
    text(s, title, 1.20, 0.30, 11.45, 0.58, 28, NAVY, True)
    rect(s, 0.48, 1.03, 12.33, 0.025, BLUE)
    text(s, source, 0.53, 7.12, 11.9, 0.22, 9, GRAY)
    return s


def card(s, x, y, w, h, heading, body, color=BLUE, body_size=16):
    rect(s, x, y, w, h, LIGHT, True, PALE)
    rect(s, x, y, w, 0.12, color)
    text(s, heading, x + 0.18, y + 0.24, w - 0.36, 0.42, 19, color, True)
    text(s, body, x + 0.18, y + 0.77, w - 0.36, h - 0.89, body_size, GRAY)


def conclusion(s, content):
    rect(s, 0.50, 6.35, 12.28, 0.58, RGBColor(252, 241, 241), True)
    text(s, "■  " + content, 0.68, 6.43, 11.95, 0.39, 16, RED, True)


def note(s, content):
    s.notes_slide.notes_text_frame.text = content


deck = Presentation(SOURCE)
original_ids = list(deck.slides._sldIdLst)
extension_ids = [original_ids[14], original_ids[15]]
for sid in original_ids:
    if sid not in extension_ids:
        deck.part.drop_rel(sid.rId)
        deck.slides._sldIdLst.remove(sid)

new_ids = []

# 1
s = base(deck, "", "目标级车路空间配准：测量模型与求解", "现代测量数据处理 · 课程小组汇报")
text(s, "以 V2I-Calib 开源代码为主体", 0.78, 1.65, 11.8, 0.72, 27, BLUE, True)
rect(s, 0.80, 2.62, 11.75, 2.25, PALE, True)
text(s, "车端目标框  →  跨端关联  →  刚体外参估计  →  误差评价", 1.08, 3.03, 11.2, 0.64, 23, NAVY, True, PP_ALIGN.CENTER)
text(s, "小组原有轨迹时差实验列为扩展分析，单独标注来源。", 1.15, 4.02, 10.9, 0.42, 18, BLUE, False, PP_ALIGN.CENTER)
text(s, "汇报人：卞浩邈    武汉大学 · 测绘遥感信息工程全国重点实验室", 0.85, 5.83, 11.7, 0.40, 15, GRAY)
note(s, "课程作业。主体论文与代码归属 Qu 等作者；末尾轨迹时差页来自原小组组会 PPT。")
new_ids.append(deck.slides._sldIdLst[-1])

# 2
s = base(deck, "01", "问题与汇报边界")
card(s, 0.58, 1.42, 3.85, 3.78, "课程问题", "不同坐标系中的车端、路侧目标框，怎样建立对应关系并估计空间变换？")
card(s, 4.73, 1.42, 3.85, 3.78, "代码主体", "V2I-Calib：目标框筛选、关联、刚体配准、外参评价。讲解 oIoU 关联路线。")
card(s, 8.88, 1.42, 3.85, 3.78, "小组扩展", "原组会分析：利用连续目标轨迹估计随时间变化的空间误差与车路时差。")
conclusion(s, "论文方法与小组扩展分开讲；课程不要求复现，所有数字均标明来源。")
new_ids.append(deck.slides._sldIdLst[-1])

# 3
s = base(deck, "01", "观测对象：两端三维目标框")
card(s, 0.60, 1.48, 5.75, 3.52, "车端观测 Bᵥ", "每帧包含若干目标：类别、三维框中心、尺寸与朝向；坐标在车端激光雷达系。")
card(s, 6.98, 1.48, 5.75, 3.52, "路侧观测 Bᵢ", "同一交通场景由路侧激光雷达观测；目标数和可见范围可能不同。")
text(s, "输入不是逐点配准的原始点云，而是检测或标注得到的目标框。", 0.72, 5.48, 11.9, 0.48, 19, BLUE, True, PP_ALIGN.CENTER)
conclusion(s, "共同目标提供跨端几何约束；遮挡与漏检使关联成为求解前提。")
new_ids.append(deck.slides._sldIdLst[-1])

# 4
s = base(deck, "02", "测量模型：求两端刚体变换", "V2I-Calib 论文 §II，式 (1)–(4)")
rect(s, 0.72, 1.45, 11.88, 1.33, PALE, True)
text(s, "pᵥ = R pᵢ + t + ε", 1.00, 1.79, 11.3, 0.62, 29, NAVY, True, PP_ALIGN.CENTER)
card(s, 0.70, 3.18, 3.73, 2.37, "未知量", "旋转矩阵 R、平移向量 t；组成 4×4 齐次外参 T。", body_size=16)
card(s, 4.80, 3.18, 3.73, 2.37, "观测量", "共同目标的三维框顶点；由目标框构成稀疏特征点集。", body_size=16)
card(s, 8.90, 3.18, 3.73, 2.37, "误差", "标注或检测噪声、遮挡、错误关联以及未处理的时差。", body_size=16)
conclusion(s, "先找到对应目标，再求 R、t，使对应几何特征尽量重合。")
new_ids.append(deck.slides._sldIdLst[-1])

# 5
s = base(deck, "02", "数据与预处理：从标注框进入算法", "上游 README：DAIR-V2X cooperative，GT 目标级配置")
card(s, 0.62, 1.38, 3.76, 3.95, "数据文件", "cooperative/data_info.json\nvehicle-side/label/lidar\ninfrastructure-side/label/virtuallidar\ncooperative/calib/lidar_i2v", body_size=14)
card(s, 4.78, 1.38, 3.76, 3.95, "预处理", "选择帧对；读取两端三维框；按类别、距离和数量筛选候选目标。", body_size=16)
card(s, 8.94, 1.38, 3.76, 3.95, "参照外参", "数据集外参用于计算旋转和平移误差；输入目标级 GT 框实验不需要原始图像和点云。", body_size=16)
conclusion(s, "课堂演示以数据结构与算法流程为主；数据集未包含在提交包。")
new_ids.append(deck.slides._sldIdLst[-1])

# 6
s = base(deck, "03", "目标关联：候选变换与 oIoU", "V2I-Calib 论文 §III-A，式 (5)–(6)")
card(s, 0.58, 1.35, 3.83, 3.84, "① 假设匹配", "从可能同类的车路目标框中取候选对应，得到候选空间变换。", body_size=16)
card(s, 4.76, 1.35, 3.83, 3.84, "② 场景评分", "把一端所有目标框变换到另一端；以整体三维 IoU（oIoU）评价场景重合程度。", body_size=16)
card(s, 8.94, 1.35, 3.83, 3.84, "③ 亲和矩阵", "组合核心亲和度与类别等辅助信息，形成跨端目标两两匹配得分。", body_size=16)
text(s, "oIoU = Σ IoU₃ᴰ(Bᵥ, T Bᵢ) / max(m, n)", 1.13, 5.56, 11.0, 0.43, 20, BLUE, True, PP_ALIGN.CENTER)
conclusion(s, "评分利用整个场景约束，减轻单个相似目标造成的歧义。")
new_ids.append(deck.slides._sldIdLst[-1])

# 7
s = base(deck, "03", "从亲和矩阵到共同目标", "V2I-Calib 论文 §III-B；上游 calib/matching/engine.py")
rect(s, 0.75, 1.56, 5.10, 3.46, PALE, True)
text(s, "跨端亲和矩阵 A", 1.11, 1.93, 4.40, 0.43, 21, NAVY, True, PP_ALIGN.CENTER)
text(s, "Aᵢⱼ = 目标 i 与目标 j 的匹配得分", 1.08, 2.70, 4.45, 0.53, 17, BLUE, False, PP_ALIGN.CENTER)
text(s, "类别 / 核心几何关系 → 得分", 1.15, 3.65, 4.3, 0.48, 16, GRAY, False, PP_ALIGN.CENTER)
rect(s, 6.72, 1.56, 5.80, 3.46, LIGHT, True, PALE)
text(s, "线性分配与筛选", 7.04, 1.93, 5.18, 0.43, 21, NAVY, True, PP_ALIGN.CENTER)
text(s, "按亲和度寻找一对一对应；\n筛除低分或不稳定的匹配。", 7.16, 2.65, 5.01, 1.18, 18, BLUE, False, PP_ALIGN.CENTER)
text(s, "输出：共同目标框对", 7.25, 4.17, 4.95, 0.40, 17, GRAY, True, PP_ALIGN.CENTER)
conclusion(s, "错误匹配会直接污染外参估计，因此关联质量是关键误差来源。")
new_ids.append(deck.slides._sldIdLst[-1])

# 8
s = base(deck, "03", "外参求解：SVD 与迭代优化", "V2I-Calib 论文 §III-C；上游 calib/solvers/svd.py")
card(s, 0.58, 1.42, 3.80, 3.80, "点集构造", "从匹配目标框提取顶点或几何特征，形成两端对应点集。", body_size=16)
card(s, 4.77, 1.42, 3.80, 3.80, "刚体最小二乘", "最小化 Σ‖pᵥ,k − (R pᵢ,k + t)‖²；由 SVD 求初始 R、t。", body_size=16)
card(s, 8.96, 1.42, 3.80, 3.80, "细化与检验", "使用匹配筛选和迭代策略优化估计，并检查结果是否稳定。", body_size=16)
text(s, "T = [ R  t ; 0  1 ]", 1.04, 5.55, 11.15, 0.45, 23, BLUE, True, PP_ALIGN.CENTER)
conclusion(s, "数据处理主链：观测筛选 → 对应关系 → 最小二乘配准 → 质量判断。")
new_ids.append(deck.slides._sldIdLst[-1])

# 9
s = base(deck, "03", "提交代码：如何对应到方法步骤", "上游仓库固定版本 21531399…；详细路径见 COURSE_README.md")
rows = [
    ("实验配置", "configs/pipeline_paper3737_oiou_gt15_thr4_top6_iter2.yaml"),
    ("运行入口", "tools/run_calibration.py"),
    ("处理主流程", "calib/pipelines/object_level.py"),
    ("目标关联", "calib/matching/engine.py"),
    ("外参求解", "calib/solvers/svd.py"),
    ("结果评价", "calib/evaluation/metrics.py"),
]
for i, (label, path) in enumerate(rows):
    y = 1.38 + i * 0.72
    rect(s, 0.76, y, 11.82, 0.61, PALE if i % 2 == 0 else LIGHT, True)
    text(s, label, 1.01, y + 0.08, 2.05, 0.43, 16, NAVY, True)
    text(s, path, 3.15, y + 0.08, 9.08, 0.43, 14, BLUE)
conclusion(s, "代码为原作者的 MIT 开源实现；本组新增课程说明与展示材料。")
new_ids.append(deck.slides._sldIdLst[-1])

# 10
s = base(deck, "04", "评价指标与论文报告结果", "数值来自 V2I-Calib 原论文 Table II（easy group）；非本组复现")
card(s, 0.59, 1.39, 3.75, 2.49, "旋转误差 RRE", "估计旋转与参考旋转的夹角；单位 °。", body_size=16)
card(s, 4.79, 1.39, 3.75, 2.49, "平移误差 RTE", "估计平移与参考平移的欧氏距离；单位 m。", body_size=16)
card(s, 8.99, 1.39, 3.75, 2.49, "成功率 / 时间", "论文按 RTE < 2 m 统计成功率，并报告运行时间。", body_size=16)
rect(s, 0.76, 4.34, 11.83, 1.32, PALE, True)
text(s, "原论文 easy group：RRE 0.68°  |  RTE 0.56 m  |  成功率 96.8%  |  时间 0.21 s", 1.02, 4.64, 11.25, 0.66, 18, NAVY, True, PP_ALIGN.CENTER)
conclusion(s, "这些是论文发表结果；本次仅提交相关代码，不声称完成复现。")
new_ids.append(deck.slides._sldIdLst[-1])

# 11
s = base(deck, "04", "方法边界：空间配准之外还有时差", "V2I-Calib 空间配准；后两页为原组会 PPT 中的小组分析")
card(s, 0.61, 1.42, 5.69, 3.77, "V2I-Calib 已覆盖", "同一时刻或已对齐帧对中的目标关联与空间外参估计。输入是三维目标框。", body_size=17)
card(s, 6.97, 1.42, 5.69, 3.77, "小组扩展所研究", "目标轨迹跨越多个时刻；空间偏差与时间偏差都可能改变两端对同一目标的位置。", body_size=17)
text(s, "时间偏差引起的位移 ≈ 目标速度 × 时间偏差", 0.95, 5.57, 11.57, 0.46, 21, BLUE, True, PP_ALIGN.CENTER)
conclusion(s, "后两页保留原组会图表，作为独立分析，不归因于 V2I-Calib 代码。")
new_ids.append(deck.slides._sldIdLst[-1])

# 12 (closing slide, moved after retained slides)
s = base(deck, "05", "总结与资料来源", "本地代码包：course_v2i_calib；课程 PPT 根据原组会材料重排")
card(s, 0.59, 1.37, 3.75, 3.50, "测量模型", "共同目标框构成跨端观测，目标是估计车路刚体空间变换。", body_size=16)
card(s, 4.79, 1.37, 3.75, 3.50, "算法链", "oIoU 场景评分、目标关联、SVD 外参求解、RRE/RTE 评价。", body_size=16)
card(s, 8.99, 1.37, 3.75, 3.50, "扩展方向", "小组原有轨迹分析提示：部署中还需考虑随时间变化的误差。", body_size=16)
text(s, "论文：arxiv.org/abs/2407.10195", 0.77, 5.18, 11.7, 0.36, 14, BLUE)
text(s, "代码：github.com/MassimoQu/v2i-calib  ·  MIT License", 0.77, 5.61, 11.7, 0.36, 14, BLUE)
conclusion(s, "本课程展示区分论文报告、开源代码与小组原有实验。")
closing_id = deck.slides._sldIdLst[-1]

# Mark the retained slides visibly and correct the scope of their claims.
for idx, sid in enumerate(extension_ids):
    slide = deck.slides[0 if idx == 0 else 1]
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text.strip() == "03":
            shape.text_frame.paragraphs[0].runs[0].text = "04"
            break
    rect(slide, 9.45, 0.12, 3.12, 0.44, PALE, True)
    text(slide, "小组原有分析 · 非上游代码结果", 9.57, 0.18, 2.88, 0.30, 10, BLUE, True)
    rect(slide, 0.48, 6.96, 12.30, 0.28, WHITE)
    text(slide, "来源：原组会 PPT《基于目标轨迹的目标级车路协同定位》；未包含在 V2I-Calib 代码中。", 0.57, 7.00, 11.98, 0.21, 9, GRAY)

# The retained slides precede the closing slide.
all_ids = list(deck.slides._sldIdLst)
for sid in all_ids:
    deck.slides._sldIdLst.remove(sid)
for sid in new_ids[:11] + extension_ids + [closing_id]:
    deck.slides._sldIdLst.append(sid)

deck.save(OUTPUT)
print(f"Saved {OUTPUT} ({len(deck.slides)} slides)")
